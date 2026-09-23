# ============================================================
# GLOXMED
# AI Medical Document Explainer
#
# Install:
#   pip install PySide6 openrouter python-dotenv pymupdf python-docx reportlab
#
# OpenRouter:
#   GloxMed automatically looks for .env beside the application.
#   If .env does not exist, it creates one automatically.
#
# Python 3.10+
# ============================================================

import os
import sys
import json
import re
from pathlib import Path
from datetime import datetime

from dotenv import load_dotenv, set_key
from openrouter import OpenRouter
import pymupdf
from docx import Document

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)
from reportlab.lib.units import mm

from PySide6.QtCore import Qt, QThread, Signal
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QApplication,
    QWidget,
    QMainWindow,
    QVBoxLayout,
    QHBoxLayout,
    QLabel,
    QPushButton,
    QTextEdit,
    QLineEdit,
    QComboBox,
    QListWidget,
    QListWidgetItem,
    QFrame,
    QFileDialog,
    QMessageBox,
    QDialog,
    QFormLayout,
    QSpinBox,
    QDoubleSpinBox,
    QAbstractItemView,
    QSizePolicy,
    QProgressBar,
)


# ============================================================
# APPLICATION PATH
# ============================================================

def get_app_dir():
    """
    IMPORTANT FOR PYINSTALLER:

    Normal Python:
        Uses the folder containing gloxmed.py

    PyInstaller one-file:
        Uses the folder containing GloxMed.exe

    This prevents .env, profiles and reports from being written
    into PyInstaller's temporary extraction directory.
    """
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent

    return Path(__file__).resolve().parent


APP_DIR = get_app_dir()

ENV_FILE = APP_DIR / ".env"
CONFIG_FILE = APP_DIR / "gloxmed_config.json"
PROFILES_FILE = APP_DIR / "gloxmed_profiles.json"
OUTPUT_DIR = APP_DIR / "GloxMed_Reports"


# Create required local folders/files
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

if not ENV_FILE.exists():
    ENV_FILE.touch()

# Load .env from BESIDE the .exe / .py file
load_dotenv(dotenv_path=ENV_FILE, override=True)


APP_NAME = "GloxMed"
APP_VERSION = "1.0"


# ============================================================
# THEME
# ============================================================

THEME = {
    "window": "#0d0e10",
    "surface": "#15171a",
    "surface2": "#1c1f23",
    "surface3": "#23272c",
    "input": "#101214",
    "text": "#f1f2f3",
    "muted": "#969ba3",
    "accent": "#d7d9dc",
    "accent_dark": "#aeb2b7",
    "border": "#30343a",
    "danger": "#d86d6d",
    "success": "#8fbd9b",
}


# ============================================================
# HELPERS
# ============================================================

def load_json(path, default):
    try:
        if path.exists():
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
    except Exception:
        pass

    return default


def save_json(path, data):
    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        with open(path, "w", encoding="utf-8") as f:
            json.dump(
                data,
                f,
                indent=4,
                ensure_ascii=False,
            )

        return True

    except Exception:
        return False


def esc(text):
    """Escape text for ReportLab Paragraph."""
    text = str(text or "")
    text = text.replace("&", "&amp;")
    text = text.replace("<", "&lt;")
    text = text.replace(">", "&gt;")
    text = text.replace("\n", "<br/>")
    return text


def clean_markdown(text):
    text = str(text or "")

    text = re.sub(
        r"```(?:json|markdown|text)?",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = text.replace("```", "")

    return text.strip()


def next_report_number():
    highest = 0

    pattern = re.compile(
        r"^(?:MediNote|MediReport)_(\d+)\.pdf$",
        re.IGNORECASE,
    )

    for p in OUTPUT_DIR.glob("*.pdf"):
        match = pattern.match(p.name)

        if match:
            highest = max(
                highest,
                int(match.group(1)),
            )

    return highest + 1


# ============================================================
# DOCUMENT EXTRACTION
# ============================================================

def extract_pdf(path):
    text_parts = []

    with pymupdf.open(str(path)) as pdf:
        for page in pdf:
            text_parts.append(
                page.get_text("text")
            )

    return "\n\n".join(text_parts).strip()


def extract_docx(path):
    doc = Document(str(path))

    parts = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            parts.append(
                paragraph.text.strip()
            )

    for table in doc.tables:
        rows = []

        for row in table.rows:
            rows.append(
                " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )
            )

        if rows:
            parts.append(
                "\n".join(rows)
            )

    return "\n\n".join(parts).strip()


def extract_document(path):
    suffix = Path(path).suffix.lower()

    if suffix == ".pdf":
        return extract_pdf(path)

    if suffix == ".docx":
        return extract_docx(path)

    raise ValueError(
        "Only PDF and DOCX files are supported."
    )


# ============================================================
# OPENROUTER
# ============================================================

def get_api_key():
    """
    Always reload the .env beside the application.

    This is important when the application is a PyInstaller
    one-file executable.
    """
    try:
        load_dotenv(
            dotenv_path=ENV_FILE,
            override=True,
        )
    except Exception:
        pass

    return os.getenv(
        "OPENROUTER_API_KEY",
        "",
    ).strip()


def save_api_key(key):
    key = key.strip()

    ENV_FILE.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not ENV_FILE.exists():
        ENV_FILE.touch()

    set_key(
        str(ENV_FILE),
        "OPENROUTER_API_KEY",
        key,
    )

    os.environ[
        "OPENROUTER_API_KEY"
    ] = key


def remove_api_key():
    try:
        if ENV_FILE.exists():
            lines = ENV_FILE.read_text(
                encoding="utf-8"
            ).splitlines()

            kept = [
                line
                for line in lines
                if not line.strip().startswith(
                    "OPENROUTER_API_KEY="
                )
            ]

            ENV_FILE.write_text(
                "\n".join(kept)
                + ("\n" if kept else ""),
                encoding="utf-8",
            )

        os.environ.pop(
            "OPENROUTER_API_KEY",
            None,
        )

        return True

    except Exception:
        return False


def get_client():
    key = get_api_key()

    if not key:
        raise RuntimeError(
            "No OpenRouter API key is configured."
        )

    return OpenRouter(
        api_key=key,
        http_referer="https://github.com/AvgLucer/GloxMed",
        x_open_router_title="GloxMed",
        x_open_router_categories="health,ai,student",
    )


def response_text(response):
    try:
        content = (
            response
            .choices[0]
            .message
            .content
        )

        if content:
            return content

    except Exception:
        pass

    try:
        return response.choices[0][
            "message"
        ]["content"]

    except Exception:
        return ""


def model_to_dict(model):
    if isinstance(model, dict):
        return model

    if hasattr(model, "model_dump"):
        try:
            return model.model_dump()
        except Exception:
            pass

    if hasattr(model, "dict"):
        try:
            return model.dict()
        except Exception:
            pass

    data = {}

    for key in (
        "id",
        "name",
        "pricing",
        "architecture",
        "context_length",
        "canonical_slug",
    ):
        if hasattr(model, key):
            data[key] = getattr(
                model,
                key,
            )

    return data


def fetch_free_models():
    client = get_client()

    try:
        result = client.models.list()

    except AttributeError:
        raise RuntimeError(
            "The installed OpenRouter SDK does not expose "
            "models.list(). Please update it with:\n\n"
            "pip install -U openrouter"
        )

    raw_models = getattr(
        result,
        "data",
        result,
    )

    if raw_models is None:
        raw_models = []

    free = []

    for item in raw_models:
        model = model_to_dict(item)

        model_id = str(
            model.get("id", "")
            or ""
        )

        name = str(
            model.get("name", "")
            or model_id
        )

        pricing = (
            model.get("pricing", {})
            or {}
        )

        if isinstance(pricing, dict):
            prompt_price = str(
                pricing.get(
                    "prompt",
                    "",
                )
            )

            completion_price = str(
                pricing.get(
                    "completion",
                    "",
                )
            )
        else:
            prompt_price = ""
            completion_price = ""

        is_free = (
            ":free" in model_id.lower()
            or (
                prompt_price
                in {
                    "0",
                    "0.0",
                    "0.000000",
                    "0.0000000",
                }
                and completion_price
                in {
                    "0",
                    "0.0",
                    "0.000000",
                    "0.0000000",
                }
            )
        )

        architecture = (
            model.get(
                "architecture",
                {},
            )
            or {}
        )

        input_modalities = architecture.get(
            "input_modalities",
            [],
        )

        if not isinstance(
            input_modalities,
            list,
        ):
            input_modalities = []

        if is_free:
            free.append(
                {
                    "id": model_id,
                    "name": name,
                    "context_length": model.get(
                        "context_length"
                    ),
                    "input_modalities": input_modalities,
                }
            )

    free.sort(
        key=lambda x: x["name"].lower()
    )

    return free


# ============================================================
# API KEY DIALOG
# ============================================================

class APIKeyDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setWindowTitle(
            "GloxMed • OpenRouter"
        )

        self.setModal(True)
        self.setFixedWidth(520)

        layout = QVBoxLayout(self)

        layout.setContentsMargins(
            24,
            24,
            24,
            24,
        )

        layout.setSpacing(14)

        title = QLabel(
            "OpenRouter API Key"
        )

        title.setObjectName(
            "DialogTitle"
        )

        subtitle = QLabel(
            "GloxMed detects OPENROUTER_API_KEY automatically "
            "from the .env file beside the application. "
            "You can also enter or replace it here."
        )

        subtitle.setWordWrap(True)
        subtitle.setObjectName(
            "DialogMuted"
        )

        self.key_input = QLineEdit()

        self.key_input.setPlaceholderText(
            "sk-or-v1-..."
        )

        self.key_input.setEchoMode(
            QLineEdit.EchoMode.Password
        )

        self.key_input.setText(
            get_api_key()
        )

        buttons = QHBoxLayout()

        remove = QPushButton("Remove")
        cancel = QPushButton("Cancel")
        save = QPushButton("Save Key")

        save.setObjectName(
            "PrimaryButton"
        )

        remove.clicked.connect(
            self.remove_key
        )

        cancel.clicked.connect(
            self.reject
        )

        save.clicked.connect(
            self.save_key
        )

        buttons.addWidget(remove)
        buttons.addStretch()
        buttons.addWidget(cancel)
        buttons.addWidget(save)

        layout.addWidget(title)
        layout.addWidget(subtitle)
        layout.addWidget(self.key_input)
        layout.addLayout(buttons)

    def save_key(self):
        key = self.key_input.text().strip()

        if not key:
            QMessageBox.warning(
                self,
                "Missing key",
                "Enter an OpenRouter API key.",
            )
            return

        try:
            save_api_key(key)
            self.accept()

        except Exception as e:
            QMessageBox.critical(
                self,
                "Could not save key",
                str(e),
            )

    def remove_key(self):
        try:
            remove_api_key()
            self.accept()

        except Exception as e:
            QMessageBox.critical(
                self,
                "Could not remove key",
                str(e),
            )


# ============================================================
# PROFILE DIALOG
# ============================================================

class ProfileDialog(QDialog):

    def __init__(
        self,
        profile=None,
        parent=None,
    ):
        super().__init__(parent)

        self.setWindowTitle(
            "Patient Profile"
        )

        self.setFixedWidth(430)

        profile = profile or {}

        layout = QVBoxLayout(self)

        layout.setContentsMargins(
            24,
            24,
            24,
            24,
        )

        layout.setSpacing(13)

        title = QLabel(
            "Patient Profile"
        )

        title.setObjectName(
            "DialogTitle"
        )

        layout.addWidget(title)

        form = QFormLayout()
        form.setSpacing(12)

        self.name = QLineEdit(
            profile.get(
                "name",
                "",
            )
        )

        self.name.setPlaceholderText(
            "Patient name"
        )

        self.age = QSpinBox()

        self.age.setRange(
            0,
            130,
        )

        self.age.setValue(
            int(
                profile.get(
                    "age",
                    0,
                )
                or 0
            )
        )

        self.height = QDoubleSpinBox()

        self.height.setRange(
            0,
            300,
        )

        self.height.setDecimals(1)

        self.height.setSuffix(
            " cm"
        )

        self.height.setValue(
            float(
                profile.get(
                    "height",
                    0,
                )
                or 0
            )
        )

        self.gender = QComboBox()

        self.gender.addItems(
            [
                "Not specified",
                "Male",
                "Female",
                "Other",
            ]
        )

        current_gender = profile.get(
            "gender",
            "Not specified",
        )

        self.gender.setCurrentText(
            current_gender
        )

        form.addRow(
            "Name",
            self.name,
        )

        form.addRow(
            "Age",
            self.age,
        )

        form.addRow(
            "Gender",
            self.gender,
        )

        form.addRow(
            "Height",
            self.height,
        )

        layout.addLayout(form)

        buttons = QHBoxLayout()

        cancel = QPushButton(
            "Cancel"
        )

        save = QPushButton(
            "Save Profile"
        )

        save.setObjectName(
            "PrimaryButton"
        )

        cancel.clicked.connect(
            self.reject
        )

        save.clicked.connect(
            self.accept
        )

        buttons.addStretch()
        buttons.addWidget(cancel)
        buttons.addWidget(save)

        layout.addLayout(buttons)

    def get_profile(self):
        return {
            "name": self.name.text().strip(),
            "age": self.age.value(),
            "gender": self.gender.currentText(),
            "height": self.height.value(),
        }


# ============================================================
# WORKER
# ============================================================

class Worker(QThread):

    finished = Signal(object)
    failed = Signal(str)

    def __init__(
        self,
        fn,
        *args,
        **kwargs,
    ):
        super().__init__()

        self.fn = fn
        self.args = args
        self.kwargs = kwargs

    def run(self):
        try:
            result = self.fn(
                *self.args,
                **self.kwargs,
            )

            self.finished.emit(
                result
            )

        except Exception as e:
            self.failed.emit(
                str(e)
            )


# ============================================================
# PDF STYLES
# ============================================================

def pdf_styles():

    styles = getSampleStyleSheet()

    return {
        "title": ParagraphStyle(
            "GloxTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=23,
            leading=27,
            textColor=colors.HexColor(
                "#17191c"
            ),
            spaceAfter=8,
        ),

        "subtitle": ParagraphStyle(
            "GloxSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor(
                "#656a71"
            ),
            spaceAfter=15,
        ),

        "section": ParagraphStyle(
            "GloxSection",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor(
                "#17191c"
            ),
            spaceBefore=12,
            spaceAfter=7,
        ),

        "body": ParagraphStyle(
            "GloxBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=9.7,
            leading=15,
            textColor=colors.HexColor(
                "#25282c"
            ),
            spaceAfter=6,
        ),

        "small": ParagraphStyle(
            "GloxSmall",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=7.8,
            leading=11,
            textColor=colors.HexColor(
                "#6d7278"
            ),
        ),

        "note": ParagraphStyle(
            "GloxNote",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor(
                "#454a50"
            ),
            backColor=colors.HexColor(
                "#f0f1f2"
            ),
            borderColor=colors.HexColor(
                "#d8dade"
            ),
            borderWidth=0.6,
            borderPadding=8,
            spaceBefore=5,
            spaceAfter=8,
        ),
    }


# ============================================================
# AI SECTION PARSER
# ============================================================

def parse_ai_sections(text):

    text = clean_markdown(text)

    lines = [
        line.strip()
        for line in text.splitlines()
    ]

    sections = []

    current_title = None
    current_lines = []

    def flush():
        nonlocal current_title
        nonlocal current_lines

        if current_title:
            body = "\n".join(
                x
                for x in current_lines
                if x
            ).strip()

            sections.append(
                (
                    current_title,
                    body,
                )
            )

        current_title = None
        current_lines = []

    for line in lines:

        if not line:
            continue

        heading = re.match(
            r"^#{1,4}\s+(.+?)\s*$",
            line,
        )

        bold_heading = re.match(
            r"^\*\*(.+?)\*\*:?\s*$",
            line,
        )

        if heading:
            flush()
            current_title = (
                heading.group(1).strip()
            )

        elif bold_heading:
            flush()
            current_title = (
                bold_heading.group(1).strip()
            )

        else:
            current_lines.append(line)

    flush()

    if not sections:
        sections = [
            (
                "AI Analysis",
                text,
            )
        ]

    return sections


# ============================================================
# PDF HEADER
# ============================================================

def make_header(
    story,
    title,
    subtitle,
):

    styles = pdf_styles()

    brand_style = ParagraphStyle(
        "brand",
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=colors.HexColor(
            "#656a71"
        ),
        spaceAfter=18,
    )

    story.append(
        Paragraph(
            "GLOXMED",
            brand_style,
        )
    )

    story.append(
        Paragraph(
            esc(title),
            styles["title"],
        )
    )

    story.append(
        Paragraph(
            esc(subtitle),
            styles["subtitle"],
        )
    )


# ============================================================
# PDF SECTIONS
# ============================================================

def add_sections(
    story,
    sections,
):

    styles = pdf_styles()

    for title, body in sections:

        story.append(
            Paragraph(
                esc(title),
                styles["section"],
            )
        )

        chunks = [
            x.strip()
            for x in body.split("\n")
            if x.strip()
        ]

        for chunk in chunks:

            if chunk.startswith(
                ("-", "•", "*")
            ):
                chunk = re.sub(
                    r"^[-•*]\s*",
                    "• ",
                    chunk,
                )

            story.append(
                Paragraph(
                    esc(chunk),
                    styles["body"],
                )
            )


# ============================================================
# MEDINOTE PDF
# ============================================================

def build_medinote(
    path,
    analysis,
    source_name,
    generated_at,
):

    styles = pdf_styles()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=17 * mm,
        bottomMargin=17 * mm,
        title="GloxMed MediNote",
        author="GloxMed",
    )

    story = []

    make_header(
        story,
        "MediNote",
        "An easy-to-understand explanation of the provided medical document.",
    )

    metadata = Table(
        [
            [
                "Source document",
                source_name,
            ],
            [
                "Generated",
                generated_at,
            ],
        ],
        colWidths=[
            38 * mm,
            135 * mm,
        ],
    )

    metadata.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#e8e9eb"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#555a60"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (1, 0),
                    (1, -1),
                    colors.HexColor(
                        "#202328"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (1, 0),
                    (1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8.2,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#d2d4d7"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(metadata)

    story.append(
        Spacer(
            1,
            9 * mm,
        )
    )

    story.append(
        Paragraph(
            esc(
                "Important: This document explains information contained "
                "in the provided document. It is not a diagnosis and does "
                "not replace a doctor's assessment."
            ),
            styles["note"],
        )
    )

    add_sections(
        story,
        parse_ai_sections(
            analysis
        ),
    )

    story.append(
        Spacer(
            1,
            10 * mm,
        )
    )

    story.append(
        Paragraph(
            "GloxMed • AI-assisted document explanation • "
            "Verify important medical decisions with a qualified "
            "healthcare professional.",
            styles["small"],
        )
    )

    doc.build(story)


# ============================================================
# MEDIREPORT PDF
# ============================================================

def build_medireport(
    path,
    analysis,
    source_name,
    generated_at,
    profile,
    instructions,
    symptoms,
):

    styles = pdf_styles()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=17 * mm,
        leftMargin=17 * mm,
        topMargin=17 * mm,
        bottomMargin=17 * mm,
        title="GloxMed MediReport",
        author="GloxMed",
    )

    story = []

    make_header(
        story,
        "MediReport",
        "Patient-context report prepared for discussion with a healthcare professional.",
    )

    profile_data = [
        [
            "Patient",
            profile.get("name")
            or "Not provided",
        ],
        [
            "Age",
            str(
                profile.get("age")
                or "Not provided"
            ),
        ],
        [
            "Gender",
            profile.get("gender")
            or "Not provided",
        ],
        [
            "Height",
            (
                f'{profile.get("height"):.1f} cm'
                if profile.get("height")
                else "Not provided"
            ),
        ],
        [
            "Weight",
            (
                f'{profile.get("weight"):.1f} kg'
                if profile.get("weight")
                else "Not provided"
            ),
        ],
        [
            "Source document",
            source_name,
        ],
        [
            "Generated",
            generated_at,
        ],
    ]

    table = Table(
        profile_data,
        colWidths=[
            38 * mm,
            135 * mm,
        ],
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#e8e9eb"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (0, -1),
                    colors.HexColor(
                        "#555a60"
                    ),
                ),
                (
                    "TEXTCOLOR",
                    (1, 0),
                    (1, -1),
                    colors.HexColor(
                        "#202328"
                    ),
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (0, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTNAME",
                    (1, 0),
                    (1, -1),
                    "Helvetica",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    8.2,
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.HexColor(
                        "#d2d4d7"
                    ),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "TOP",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    7,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(table)

    story.append(
        Spacer(
            1,
            7 * mm,
        )
    )

    if instructions.strip():

        story.append(
            Paragraph(
                "Doctor's Instructions / Advice",
                styles["section"],
            )
        )

        story.append(
            Paragraph(
                esc(instructions),
                styles["body"],
            )
        )

    if symptoms.strip():

        story.append(
            Paragraph(
                "What the Patient Experienced",
                styles["section"],
            )
        )

        story.append(
            Paragraph(
                esc(symptoms),
                styles["body"],
            )
        )

    story.append(
        Paragraph(
            esc(
                "Clinical-use note: This report is a communication aid. "
                "The patient profile and reported symptoms are user-provided "
                "and should be confirmed by the treating healthcare professional."
            ),
            styles["note"],
        )
    )

    story.append(
        PageBreak()
    )

    story.append(
        Paragraph(
            "AI-Assisted Explanation",
            styles["section"],
        )
    )

    add_sections(
        story,
        parse_ai_sections(
            analysis
        ),
    )

    story.append(
        Spacer(
            1,
            10 * mm,
        )
    )

    story.append(
        Paragraph(
            "GloxMed • Bring this report together with the original "
            "prescription/report when consulting a healthcare professional.",
            styles["small"],
        )
    )

    doc.build(story)


# ============================================================
# AI PROMPT
# ============================================================

SYSTEM_PROMPT = """
You are GloxMed, a careful medical-document explanation assistant.

Your job is to read a medical document supplied by a patient and produce
a clear, accurate, non-diagnostic explanation.

CRITICAL RULES:

1. Do not diagnose the patient.

2. Do not invent laboratory values, medicines, conditions, symptoms,
   instructions, or findings that are not supported by the source.

3. Distinguish clearly between:
   - information explicitly present in the document,
   - reasonable general explanations,
   - points that should be discussed with a healthcare professional.

4. If a value, term, medicine, dosage, or instruction is unclear,
   explicitly say it is unclear rather than guessing.

5. Never tell the patient to start, stop, increase, decrease, or replace
   a medicine.

6. Never present potential suggestions as prescriptions.

7. If the source contains an existing doctor's recommendation,
   explain it faithfully.

8. Use plain, understandable English while preserving important medical terms.

9. Do not expose hidden reasoning.

10. The result will be placed into a PDF that may be shown to a doctor.

Return ONLY a clean Markdown-style report with these sections:

# Plain-English Summary

Explain the main purpose and important findings in simple language.

# Important Findings

List the clinically relevant information explicitly present in the source.

Include values, units, reference ranges, medicine names, dosage information,
or other important details when available.

# What These Terms Mean

Explain important medical terminology from the document in accessible language.

# Doctor's Advice / Instructions

Explain instructions or recommendations that are actually present in the source.
If none are present, say so.

# Potential Points to Discuss With a Doctor

Give cautious, non-prescriptive discussion points based only on the document
and the patient's supplied context. Do not invent diagnoses or treatments.

# Information That May Need Clarification

Mention ambiguous, unreadable, missing, or unclear information.

# Safety Note

End with a short reminder that this explanation does not replace professional
medical advice and that urgent symptoms should be assessed by appropriate
medical services.
"""


def generate_report_with_ai(
    model_id,
    document_text,
    instructions,
    symptoms,
):

    client = get_client()

    prompt = f"""
DOCUMENT CONTENT
----------------

{document_text}

PATIENT/USER SUPPLIED CONTEXT
----------------

Doctor's instructions/advice:

{instructions or "Not provided."}

What the patient says they actually felt:

{symptoms or "Not provided."}

Generate the GloxMed report according to your system instructions.
"""

    response = client.chat.send(
        model=model_id,
        messages=[
            {
                "role": "system",
                "content": SYSTEM_PROMPT,
            },
            {
                "role": "user",
                "content": prompt,
            },
        ],
        temperature=0.15,
        max_tokens=7000,
    )

    result = response_text(
        response
    )

    if not result.strip():
        raise RuntimeError(
            "The selected model returned an empty response."
        )

    return clean_markdown(
        result
    )


# ============================================================
# MAIN WINDOW
# ============================================================

class GloxMed(QMainWindow):

    def __init__(self):

        super().__init__()

        self.setWindowTitle(
            "GloxMed"
        )

        self.setMinimumSize(
            1180,
            760,
        )

        self.resize(
            1280,
            820,
        )

        self.profiles = load_json(
            PROFILES_FILE,
            {},
        )

        self.current_profile = None
        self.selected_file = None
        self.selected_file_text = ""
        self.worker = None
        self.drag_position = None

        self.build_ui()
        self.apply_theme()
        self.refresh_profiles()
        self.update_api_status()

        if get_api_key():
            self.load_models()

    # ========================================================
    # UI
    # ========================================================

    def build_ui(self):

        root = QWidget()

        root_layout = QVBoxLayout(
            root
        )

        root_layout.setContentsMargins(
            18,
            15,
            18,
            18,
        )

        root_layout.setSpacing(12)

        # HEADER

        header = QHBoxLayout()

        logo = QLabel(
            "GloxMed"
        )

        logo.setObjectName(
            "Logo"
        )

        logo.setFont(
            QFont(
                "Segoe UI",
                18,
                QFont.Weight.Bold,
            )
        )

        subtitle = QLabel(
            "AI-assisted medical document understanding"
        )

        subtitle.setObjectName(
            "Subtitle"
        )

        header.addWidget(logo)
        header.addSpacing(10)
        header.addWidget(subtitle)
        header.addStretch()

        self.api_status = QLabel(
            "○ No OpenRouter API key"
        )

        self.api_status.setObjectName(
            "Status"
        )

        api_btn = QPushButton(
            "API"
        )

        api_btn.clicked.connect(
            self.open_api_dialog
        )

        refresh_btn = QPushButton(
            "↻"
        )

        refresh_btn.setFixedSize(
            38,
            34,
        )

        refresh_btn.clicked.connect(
            self.load_models
        )

        minimize_btn = QPushButton(
            "—"
        )

        minimize_btn.setFixedSize(
            38,
            34,
        )

        minimize_btn.clicked.connect(
            self.showMinimized
        )

        close_btn = QPushButton(
            "×"
        )

        close_btn.setFixedSize(
            38,
            34,
        )

        close_btn.clicked.connect(
            self.close
        )

        header.addWidget(
            self.api_status
        )

        header.addSpacing(8)
        header.addWidget(api_btn)
        header.addWidget(refresh_btn)
        header.addWidget(minimize_btn)
        header.addWidget(close_btn)

        root_layout.addLayout(
            header
        )

        # MAIN TWO COLUMN AREA

        columns = QHBoxLayout()
        columns.setSpacing(12)

        # LEFT COLUMN

        left = QFrame()
        left.setObjectName(
            "Panel"
        )

        left_layout = QVBoxLayout(
            left
        )

        left_layout.setContentsMargins(
            17,
            17,
            17,
            17,
        )

        left_layout.setSpacing(11)

        left_title = QLabel(
            "MEDICAL DOCUMENT"
        )

        left_title.setObjectName(
            "SectionTitle"
        )

        left_layout.addWidget(
            left_title
        )

        self.file_card = QFrame()

        self.file_card.setObjectName(
            "FileCard"
        )

        file_layout = QVBoxLayout(
            self.file_card
        )

        file_layout.setContentsMargins(
            14,
            14,
            14,
            14,
        )

        file_layout.setSpacing(5)

        self.file_name = QLabel(
            "No document selected"
        )

        self.file_name.setObjectName(
            "FileName"
        )

        self.file_name.setWordWrap(
            True
        )

        self.file_status = QLabel(
            "Select a PDF or DOCX document."
        )

        self.file_status.setObjectName(
            "Muted"
        )

        file_layout.addWidget(
            self.file_name
        )

        file_layout.addWidget(
            self.file_status
        )

        left_layout.addWidget(
            self.file_card
        )

        select_btn = QPushButton(
            "＋  Select Medical Document"
        )

        select_btn.setObjectName(
            "PrimaryButton"
        )

        select_btn.setMinimumHeight(
            44
        )

        select_btn.clicked.connect(
            self.select_document
        )

        left_layout.addWidget(
            select_btn
        )

        instruction_label = QLabel(
            "DOCTOR'S INSTRUCTIONS / ADVICE"
        )

        instruction_label.setObjectName(
            "SectionTitle"
        )

        left_layout.addWidget(
            instruction_label
        )

        self.instructions = QTextEdit()

        self.instructions.setPlaceholderText(
            "Optional: enter what the doctor instructed the patient to do..."
        )

        self.instructions.setMaximumHeight(
            105
        )

        left_layout.addWidget(
            self.instructions
        )

        symptoms_label = QLabel(
            "WHAT DID THE PATIENT ACTUALLY FEEL?"
        )

        symptoms_label.setObjectName(
            "SectionTitle"
        )

        left_layout.addWidget(
            symptoms_label
        )

        self.symptoms = QTextEdit()

        self.symptoms.setPlaceholderText(
            "Optional: describe what the patient actually experienced..."
        )

        self.symptoms.setMaximumHeight(
            125
        )

        left_layout.addWidget(
            self.symptoms
        )

        model_row = QHBoxLayout()

        model_label = QLabel(
            "AI MODEL"
        )

        model_label.setObjectName(
            "ControlLabel"
        )

        self.model_combo = QComboBox()

        self.model_combo.setMinimumHeight(
            36
        )

        self.model_combo.setSizePolicy(
            QSizePolicy.Policy.Expanding,
            QSizePolicy.Policy.Fixed,
        )

        model_row.addWidget(
            model_label
        )

        model_row.addWidget(
            self.model_combo
        )

        left_layout.addLayout(
            model_row
        )

        self.generate_btn = QPushButton(
            "✦  GENERATE MEDINOTE + MEDIREPORT"
        )

        self.generate_btn.setObjectName(
            "GenerateButton"
        )

        self.generate_btn.setMinimumHeight(
            48
        )

        self.generate_btn.clicked.connect(
            self.generate_reports
        )

        left_layout.addWidget(
            self.generate_btn
        )

        self.progress = QProgressBar()

        self.progress.setRange(
            0,
            0,
        )

        self.progress.setTextVisible(
            False
        )

        self.progress.hide()

        left_layout.addWidget(
            self.progress
        )

        self.document_preview = QLabel(
            "The selected document's text will be processed locally "
            "and only the extracted text plus the supplied context is "
            "sent to the selected OpenRouter model."
        )

        self.document_preview.setObjectName(
            "PrivacyNote"
        )

        self.document_preview.setWordWrap(
            True
        )

        left_layout.addWidget(
            self.document_preview
        )

        left_layout.addStretch()

        # RIGHT COLUMN

        right = QFrame()

        right.setObjectName(
            "Panel"
        )

        right_layout = QVBoxLayout(
            right
        )

        right_layout.setContentsMargins(
            17,
            17,
            17,
            17,
        )

        right_layout.setSpacing(11)

        profile_header = QHBoxLayout()

        profile_title = QLabel(
            "PATIENT PROFILES"
        )

        profile_title.setObjectName(
            "SectionTitle"
        )

        add_profile = QPushButton(
            "＋"
        )

        add_profile.setFixedSize(
            34,
            32,
        )

        add_profile.setToolTip(
            "Create profile"
        )

        add_profile.clicked.connect(
            self.add_profile
        )

        edit_profile = QPushButton(
            "Edit"
        )

        edit_profile.clicked.connect(
            self.edit_profile
        )

        delete_profile = QPushButton(
            "Delete"
        )

        delete_profile.clicked.connect(
            self.delete_profile
        )

        profile_header.addWidget(
            profile_title
        )

        profile_header.addStretch()

        profile_header.addWidget(
            add_profile
        )

        profile_header.addWidget(
            edit_profile
        )

        profile_header.addWidget(
            delete_profile
        )

        right_layout.addLayout(
            profile_header
        )

        self.profile_list = QListWidget()

        self.profile_list.setSelectionMode(
            QAbstractItemView.SelectionMode.SingleSelection
        )

        self.profile_list.currentItemChanged.connect(
            self.profile_changed
        )

        right_layout.addWidget(
            self.profile_list,
            1,
        )

        details_title = QLabel(
            "CURRENT PATIENT"
        )

        details_title.setObjectName(
            "SectionTitle"
        )

        right_layout.addWidget(
            details_title
        )

        self.profile_card = QFrame()

        self.profile_card.setObjectName(
            "ProfileCard"
        )

        card_layout = QVBoxLayout(
            self.profile_card
        )

        card_layout.setContentsMargins(
            15,
            15,
            15,
            15,
        )

        card_layout.setSpacing(9)

        self.profile_name = QLabel(
            "No profile selected"
        )

        self.profile_name.setObjectName(
            "ProfileName"
        )

        self.profile_details = QLabel(
            "Create a patient profile to include patient details in MediReport."
        )

        self.profile_details.setObjectName(
            "Muted"
        )

        self.profile_details.setWordWrap(
            True
        )

        card_layout.addWidget(
            self.profile_name
        )

        card_layout.addWidget(
            self.profile_details
        )

        right_layout.addWidget(
            self.profile_card
        )

        weight_title = QLabel(
            "WEIGHT FOR THIS REPORT"
        )

        weight_title.setObjectName(
            "SectionTitle"
        )

        right_layout.addWidget(
            weight_title
        )

        self.weight = QDoubleSpinBox()

        self.weight.setRange(
            0,
            500,
        )

        self.weight.setDecimals(
            1
        )

        self.weight.setSuffix(
            " kg"
        )

        self.weight.setSpecialValueText(
            "Not provided"
        )

        self.weight.setMinimumHeight(
            36
        )

        right_layout.addWidget(
            self.weight
        )

        report_info = QLabel(
            "MediNote_XX.pdf\n"
            "Plain-English explanation\n\n"
            "MediReport_XX.pdf\n"
            "Patient profile + context + explanation"
        )

        report_info.setObjectName(
            "OutputInfo"
        )

        right_layout.addWidget(
            report_info
        )

        self.last_output = QLabel(
            "Reports will appear in:\n"
            + str(OUTPUT_DIR)
        )

        self.last_output.setObjectName(
            "Muted"
        )

        self.last_output.setWordWrap(
            True
        )

        right_layout.addWidget(
            self.last_output
        )

        columns.addWidget(
            left,
            3,
        )

        columns.addWidget(
            right,
            2,
        )

        root_layout.addLayout(
            columns,
            1,
        )

        footer = QLabel(
            "GloxMed  •  OpenRouter powered  •  "
            "Keep the original medical document with the generated report"
        )

        footer.setObjectName(
            "Footer"
        )

        footer.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        root_layout.addWidget(
            footer
        )

        self.setCentralWidget(
            root
        )

    # ========================================================
    # THEME
    # ========================================================

    def apply_theme(self):

        t = THEME

        self.setStyleSheet(
            f"""
            QWidget {{
                font-family: "Segoe UI";
                color: {t["text"]};
            }}

            QMainWindow {{
                background: {t["window"]};
            }}

            QFrame#Panel {{
                background: {t["surface"]};
                border: 1px solid {t["border"]};
                border-radius: 18px;
            }}

            QFrame#FileCard,
            QFrame#ProfileCard {{
                background: {t["surface2"]};
                border: 1px solid {t["border"]};
                border-radius: 13px;
            }}

            QLabel#Logo {{
                color: {t["text"]};
            }}

            QLabel#Subtitle {{
                color: {t["muted"]};
                font-size: 11px;
            }}

            QLabel#SectionTitle {{
                color: {t["muted"]};
                font-size: 10px;
                font-weight: 700;
                letter-spacing: 1px;
            }}

            QLabel#ControlLabel {{
                color: {t["muted"]};
                font-size: 10px;
                font-weight: 700;
            }}

            QLabel#Status {{
                color: {t["accent"]};
                font-size: 10px;
            }}

            QLabel#FileName {{
                color: {t["text"]};
                font-size: 13px;
                font-weight: 700;
            }}

            QLabel#ProfileName {{
                color: {t["text"]};
                font-size: 15px;
                font-weight: 700;
            }}

            QLabel#Muted,
            QLabel#PrivacyNote,
            QLabel#OutputInfo {{
                color: {t["muted"]};
                font-size: 10px;
            }}

            QLabel#PrivacyNote {{
                background: {t["input"]};
                border: 1px solid {t["border"]};
                border-radius: 10px;
                padding: 9px;
            }}

            QLabel#Footer {{
                color: {t["muted"]};
                font-size: 9px;
            }}

            QTextEdit,
            QLineEdit,
            QSpinBox,
            QDoubleSpinBox,
            QComboBox {{
                background: {t["input"]};
                color: {t["text"]};
                border: 1px solid {t["border"]};
                border-radius: 10px;
                padding: 8px 10px;
                selection-background-color: {t["surface3"]};
            }}

            QTextEdit:focus,
            QLineEdit:focus,
            QSpinBox:focus,
            QDoubleSpinBox:focus,
            QComboBox:focus {{
                border: 1px solid {t["accent_dark"]};
            }}

            QComboBox QAbstractItemView {{
                background: {t["surface2"]};
                color: {t["text"]};
                border: 1px solid {t["border"]};
                selection-background-color: {t["surface3"]};
            }}

            QListWidget {{
                background: {t["input"]};
                color: {t["text"]};
                border: 1px solid {t["border"]};
                border-radius: 12px;
                padding: 6px;
                outline: none;
            }}

            QListWidget::item {{
                padding: 10px;
                border-radius: 9px;
                margin: 2px;
            }}

            QListWidget::item:selected {{
                background: {t["surface3"]};
                color: {t["text"]};
                border: 1px solid {t["border"]};
            }}

            QPushButton {{
                background: {t["surface2"]};
                color: {t["text"]};
                border: 1px solid {t["border"]};
                border-radius: 9px;
                padding: 8px 13px;
                font-weight: 600;
            }}

            QPushButton:hover {{
                background: {t["surface3"]};
                border: 1px solid {t["accent_dark"]};
            }}

            QPushButton#PrimaryButton {{
                background: {t["surface3"]};
                color: {t["text"]};
                border: 1px solid {t["accent_dark"]};
                font-weight: 700;
            }}

            QPushButton#GenerateButton {{
                background: {t["accent"]};
                color: #111214;
                border: none;
                border-radius: 11px;
                font-weight: 800;
                font-size: 11px;
            }}

            QPushButton#GenerateButton:hover {{
                background: #ffffff;
            }}

            QProgressBar {{
                background: {t["input"]};
                border: 1px solid {t["border"]};
                border-radius: 5px;
                height: 7px;
            }}

            QProgressBar::chunk {{
                background: {t["accent"]};
                border-radius: 5px;
            }}

            QScrollBar:vertical {{
                background: transparent;
                width: 7px;
            }}

            QScrollBar::handle:vertical {{
                background: {t["border"]};
                border-radius: 4px;
            }}

            QDialog {{
                background: {t["surface"]};
                color: {t["text"]};
            }}

            QLabel#DialogTitle {{
                color: {t["text"]};
                font-size: 16px;
                font-weight: 800;
            }}

            QLabel#DialogMuted {{
                color: {t["muted"]};
                font-size: 10px;
            }}
            """
        )

    # ========================================================
    # API
    # ========================================================

    def update_api_status(self):

        if get_api_key():

            self.api_status.setText(
                "● OpenRouter connected"
            )

            self.api_status.setStyleSheet(
                f"color: {THEME['success']};"
            )

        else:

            self.api_status.setText(
                "○ No OpenRouter API key"
            )

            self.api_status.setStyleSheet(
                f"color: {THEME['muted']};"
            )

    def open_api_dialog(self):

        dialog = APIKeyDialog(
            self
        )

        if (
            dialog.exec()
            == QDialog.DialogCode.Accepted
        ):

            load_dotenv(
                dotenv_path=ENV_FILE,
                override=True,
            )

            self.update_api_status()

            if get_api_key():
                self.load_models()

            else:
                self.model_combo.clear()

                self.model_combo.addItem(
                    "Add an OpenRouter API key"
                )

    def load_models(self):

        if not get_api_key():

            self.model_combo.clear()

            self.model_combo.addItem(
                "Add an OpenRouter API key"
            )

            return

        self.model_combo.clear()

        self.model_combo.addItem(
            "Detecting free OpenRouter models..."
        )

        self.model_combo.setEnabled(
            False
        )

        self.worker = Worker(
            fetch_free_models
        )

        self.worker.finished.connect(
            self.models_loaded
        )

        self.worker.failed.connect(
            self.models_failed
        )

        self.worker.start()

    def models_loaded(
        self,
        models,
    ):

        self.model_combo.clear()

        if not models:

            self.model_combo.addItem(
                "No free models found"
            )

            self.model_combo.setEnabled(
                False
            )

            return

        self.model_combo.setEnabled(
            True
        )

        for model in models:

            label = model["name"]

            model_id = model["id"]

            if model.get(
                "context_length"
            ):

                label += (
                    f"  •  "
                    f"{model['context_length']:,} ctx"
                )

            self.model_combo.addItem(
                label,
                model_id,
            )

        preferred_terms = [
            "nemotron",
            "gpt-oss",
            "qwen",
            "deepseek",
            "gemma",
            "llama",
        ]

        for term in preferred_terms:

            for i in range(
                self.model_combo.count()
            ):

                model_id = (
                    self.model_combo.itemData(i)
                )

                if (
                    model_id
                    and term
                    in model_id.lower()
                ):

                    self.model_combo.setCurrentIndex(
                        i
                    )

                    return

    def models_failed(
        self,
        message,
    ):

        self.model_combo.clear()

        self.model_combo.addItem(
            "Could not load free models"
        )

        self.model_combo.setEnabled(
            False
        )

        self.api_status.setText(
            "● API key detected • model list failed"
        )

        QMessageBox.warning(
            self,
            "Model detection failed",
            f"{message}\n\n"
            "Try refreshing after checking your OpenRouter key.",
        )

    # ========================================================
    # DOCUMENT
    # ========================================================

    def select_document(self):

        path, _ = QFileDialog.getOpenFileName(
            self,
            "Select Medical Document",
            str(Path.home()),
            "Medical Documents (*.pdf *.docx);;"
            "PDF (*.pdf);;"
            "DOCX (*.docx)",
        )

        if not path:
            return

        try:

            text = extract_document(
                path
            )

            if not text.strip():

                raise ValueError(
                    "No selectable text was found in this document. "
                    "A scanned/image-only document needs OCR before "
                    "GloxMed can read it."
                )

            self.selected_file = path
            self.selected_file_text = text

            self.file_name.setText(
                Path(path).name
            )

            self.file_status.setText(
                f"Ready • {len(text):,} extracted characters"
            )

            self.document_preview.setText(
                "Document text extracted locally. "
                "GloxMed will send the extracted text, doctor instructions, "
                "and patient-reported symptoms to the selected model."
            )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Could not read document",
                str(e),
            )

    # ========================================================
    # PROFILES
    # ========================================================

    def refresh_profiles(self):

        self.profile_list.blockSignals(
            True
        )

        self.profile_list.clear()

        for (
            profile_id,
            profile,
        ) in self.profiles.items():

            item = QListWidgetItem(
                profile.get(
                    "name"
                )
                or "Unnamed patient"
            )

            item.setData(
                Qt.ItemDataRole.UserRole,
                profile_id,
            )

            self.profile_list.addItem(
                item
            )

        self.profile_list.blockSignals(
            False
        )

        if (
            self.current_profile
            and self.current_profile
            in self.profiles
        ):

            for i in range(
                self.profile_list.count()
            ):

                item = (
                    self.profile_list.item(i)
                )

                if (
                    item.data(
                        Qt.ItemDataRole.UserRole
                    )
                    == self.current_profile
                ):

                    self.profile_list.setCurrentRow(
                        i
                    )

                    return

        if self.profile_list.count():

            self.profile_list.setCurrentRow(
                0
            )

        else:

            self.current_profile = None
            self.update_profile_card()

    def profile_changed(
        self,
        current,
        previous,
    ):

        if not current:

            self.current_profile = None
            self.update_profile_card()

            return

        profile_id = current.data(
            Qt.ItemDataRole.UserRole
        )

        self.current_profile = (
            profile_id
        )

        self.update_profile_card()

    def update_profile_card(self):

        profile = self.profiles.get(
            self.current_profile
        )

        if not profile:

            self.profile_name.setText(
                "No profile selected"
            )

            self.profile_details.setText(
                "Create a patient profile to include patient details in MediReport."
            )

            return

        self.profile_name.setText(
            profile.get(
                "name"
            )
            or "Unnamed patient"
        )

        height = profile.get(
            "height",
            0,
        )

        age = profile.get(
            "age",
            0,
        )

        gender = profile.get(
            "gender",
            "Not specified",
        )

        height_text = (
            f"{height:.1f} cm"
            if height
            else "Height not set"
        )

        self.profile_details.setText(
            f"Age: {age or 'Not set'}  •  "
            f"Gender: {gender}  •  "
            f"Height: {height_text}"
        )

    def add_profile(self):

        dialog = ProfileDialog(
            parent=self
        )

        if (
            dialog.exec()
            != QDialog.DialogCode.Accepted
        ):
            return

        profile = dialog.get_profile()

        if not profile["name"]:

            QMessageBox.warning(
                self,
                "Name required",
                "Give the patient profile a name.",
            )

            return

        profile_id = (
            "profile_"
            + datetime.now().strftime(
                "%Y%m%d%H%M%S%f"
            )
        )

        self.profiles[
            profile_id
        ] = profile

        self.current_profile = (
            profile_id
        )

        save_json(
            PROFILES_FILE,
            self.profiles,
        )

        self.refresh_profiles()

    def edit_profile(self):

        if not self.current_profile:

            QMessageBox.information(
                self,
                "No profile",
                "Select a patient profile first.",
            )

            return

        profile = self.profiles.get(
            self.current_profile
        )

        if not profile:
            return

        dialog = ProfileDialog(
            profile,
            self,
        )

        if (
            dialog.exec()
            != QDialog.DialogCode.Accepted
        ):
            return

        updated = (
            dialog.get_profile()
        )

        if not updated["name"]:

            QMessageBox.warning(
                self,
                "Name required",
                "Give the patient profile a name.",
            )

            return

        self.profiles[
            self.current_profile
        ] = updated

        save_json(
            PROFILES_FILE,
            self.profiles,
        )

        self.refresh_profiles()

    def delete_profile(self):

        if not self.current_profile:
            return

        profile = self.profiles.get(
            self.current_profile,
            {},
        )

        name = profile.get(
            "name",
            "this profile",
        )

        answer = QMessageBox.question(
            self,
            "Delete profile",
            f"Delete {name}'s profile?",
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No,
        )

        if (
            answer
            != QMessageBox.StandardButton.Yes
        ):
            return

        self.profiles.pop(
            self.current_profile,
            None,
        )

        self.current_profile = None

        save_json(
            PROFILES_FILE,
            self.profiles,
        )

        self.refresh_profiles()

    # ========================================================
    # GENERATION
    # ========================================================

    def generate_reports(self):

        if not get_api_key():

            answer = QMessageBox.question(
                self,
                "OpenRouter key required",
                "No OpenRouter API key was detected. "
                "Open API settings now?",
                QMessageBox.StandardButton.Yes
                | QMessageBox.StandardButton.No,
            )

            if (
                answer
                == QMessageBox.StandardButton.Yes
            ):
                self.open_api_dialog()

            return

        if (
            not self.selected_file
            or not self.selected_file_text
        ):

            QMessageBox.warning(
                self,
                "No document",
                "Select a PDF or DOCX medical document first.",
            )

            return

        model_id = (
            self.model_combo.currentData()
        )

        if not model_id:

            QMessageBox.warning(
                self,
                "No model",
                "Select a free OpenRouter model.",
            )

            return

        self.set_busy(
            True
        )

        instructions = (
            self.instructions
            .toPlainText()
            .strip()
        )

        symptoms = (
            self.symptoms
            .toPlainText()
            .strip()
        )

        args = (
            model_id,
            self.selected_file_text,
            instructions,
            symptoms,
        )

        self.worker = Worker(
            generate_report_with_ai,
            *args,
        )

        self.worker.finished.connect(
            self.ai_finished
        )

        self.worker.failed.connect(
            self.ai_failed
        )

        self.worker.start()

    def set_busy(
        self,
        busy,
    ):

        self.generate_btn.setEnabled(
            not busy
        )

        if busy:

            self.generate_btn.setText(
                "✦  ANALYZING MEDICAL DOCUMENT..."
            )

            self.progress.show()

        else:

            self.generate_btn.setText(
                "✦  GENERATE MEDINOTE + MEDIREPORT"
            )

            self.progress.hide()

    def ai_finished(
        self,
        analysis,
    ):

        try:

            number = next_report_number()

            note_path = (
                OUTPUT_DIR
                / f"MediNote_{number:02d}.pdf"
            )

            report_path = (
                OUTPUT_DIR
                / f"MediReport_{number:02d}.pdf"
            )

            generated_at = (
                datetime.now().strftime(
                    "%d %B %Y, %I:%M %p"
                )
            )

            profile = dict(
                self.profiles.get(
                    self.current_profile,
                    {},
                )
            )

            profile["weight"] = (
                self.weight.value()
            )

            build_medinote(
                note_path,
                analysis,
                Path(
                    self.selected_file
                ).name,
                generated_at,
            )

            build_medireport(
                report_path,
                analysis,
                Path(
                    self.selected_file
                ).name,
                generated_at,
                profile,
                self.instructions
                .toPlainText()
                .strip(),
                self.symptoms
                .toPlainText()
                .strip(),
            )

            self.last_output.setText(
                "Generated successfully:\n\n"
                f"{note_path.name}\n"
                f"{report_path.name}\n\n"
                "Folder:\n"
                f"{OUTPUT_DIR}"
            )

            self.set_busy(
                False
            )

            answer = QMessageBox.question(
                self,
                "GloxMed reports ready",
                f"Created:\n\n"
                f"{note_path.name}\n"
                f"{report_path.name}\n\n"
                "Open the output folder?",
                QMessageBox.StandardButton.Yes
                | QMessageBox.StandardButton.No,
            )

            if (
                answer
                == QMessageBox.StandardButton.Yes
            ):
                self.open_output_folder()

        except Exception as e:

            self.set_busy(
                False
            )

            QMessageBox.critical(
                self,
                "PDF generation failed",
                str(e),
            )

    def ai_failed(
        self,
        message,
    ):

        self.set_busy(
            False
        )

        QMessageBox.critical(
            self,
            "OpenRouter generation failed",
            message,
        )

    # ========================================================
    # OUTPUT
    # ========================================================

    def open_output_folder(self):

        try:

            if sys.platform.startswith(
                "win"
            ):

                os.startfile(
                    str(OUTPUT_DIR)
                )

            elif sys.platform == "darwin":

                os.system(
                    f'open "{OUTPUT_DIR}"'
                )

            else:

                os.system(
                    f'xdg-open "{OUTPUT_DIR}"'
                )

        except Exception:
            pass

    # ========================================================
    # WINDOW
    # ========================================================

    def mousePressEvent(
        self,
        event,
    ):

        if (
            event.button()
            == Qt.MouseButton.LeftButton
        ):

            self.drag_position = (
                event.globalPosition()
                .toPoint()
                - self.frameGeometry()
                .topLeft()
            )

            event.accept()

    def mouseMoveEvent(
        self,
        event,
    ):

        if (
            event.buttons()
            & Qt.MouseButton.LeftButton
            and self.drag_position
        ):

            self.move(
                event.globalPosition()
                .toPoint()
                - self.drag_position
            )

            event.accept()

    def mouseReleaseEvent(
        self,
        event,
    ):

        self.drag_position = None

    def closeEvent(
        self,
        event,
    ):

        save_json(
            PROFILES_FILE,
            self.profiles,
        )

        event.accept()


# ============================================================
# MAIN
# ============================================================

def main():

    # Make sure these exist even in a packaged application.
    APP_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not ENV_FILE.exists():
        ENV_FILE.touch()

    load_dotenv(
        dotenv_path=ENV_FILE,
        override=True,
    )

    app = QApplication(
        sys.argv
    )

    app.setApplicationName(
        "GloxMed"
    )

    app.setApplicationVersion(
        APP_VERSION
    )

    app.setStyle(
        "Fusion"
    )

    window = GloxMed()

    screen = app.primaryScreen()

    if screen:

        geometry = (
            screen.availableGeometry()
        )

        window.move(
            geometry.center()
            - window.rect().center()
        )

    window.show()

    sys.exit(
        app.exec()
    )


if __name__ == "__main__":
    main()